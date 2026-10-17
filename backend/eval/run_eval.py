import argparse
import asyncio
import hashlib
import json
import re
import statistics
import sys
import time
from dataclasses import asdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook

from app import config
from app.analyzers.common import Event
from app.analyzers import probe, twopass
from app.pipeline import decide


def norm(name: str) -> str:
    return Path(name).stem.strip().replace("Х", "X")


def parse_time(text) -> float | None:
    if text is None:
        return None
    s = str(text)
    m = re.search(r"(\d+)\s*мин", s)
    sec = re.search(r"(\d+)\s*сек", s)
    total = (int(m.group(1)) * 60 if m else 0) + (int(sec.group(1)) if sec else 0)
    return float(total) if (m or sec) else None


def load_labels(path: Path) -> dict[str, dict]:
    ws = load_workbook(path, read_only=True).active
    out = {}
    for name, ts, label in list(ws.iter_rows(values_only=True))[1:]:
        if not name:
            continue
        out[norm(name)] = {"violation": str(label).strip().startswith("нарушение"), "t": parse_time(ts)}
    return out


def md5(path: Path) -> str:
    return hashlib.md5(path.read_bytes()).hexdigest()


def build_items(videos: Path, labels: dict) -> list[dict]:
    items: dict[str, dict] = {}
    for v in sorted(videos.glob("*.mov"), key=lambda p: int(re.sub(r"\D", "", p.stem) or 0)):
        key = norm(v.name)
        if key not in labels:
            continue
        h = md5(v)
        lab = labels[key]
        item = items.setdefault(h, {"name": key, "aliases": [], "path": v, "violation": False, "times": []})
        item["aliases"].append(key)
        item["violation"] |= lab["violation"]
        if lab["t"] is not None:
            item["times"].append(lab["t"])
    return list(items.values())


async def analyze_all(items: list[dict], out: Path, force: bool, parallel: int, candidate: float,
                      analyzer: str, scan_from: Path | None) -> None:
    sem = asyncio.Semaphore(parallel)

    async def one(item):
        d = out / item["name"]
        raw = d / "raw_events.json"
        if raw.exists() and not force:
            return
        async with sem:
            t0 = time.time()
            d.mkdir(parents=True, exist_ok=True)
            if analyzer == "twopass":
                await twopass.analyze(item["path"], d, candidate=candidate, scan_dir=scan_from / item["name"] if scan_from else None)
            else:
                await probe.analyze(item["path"], d, candidate=candidate)
            (d / "elapsed.txt").write_text(f"{time.time() - t0:.1f}")
            print(f"{item['name']}: {time.time() - t0:.1f}s", flush=True)

    await asyncio.gather(*(one(i) for i in items))


def load_events(d: Path) -> list[Event]:
    return [Event(**e) for e in json.loads((d / "raw_events.json").read_text(encoding="utf-8"))]


def score(items: list[dict], out: Path, threshold: float) -> dict:
    tp = fp = fn = tn = 0
    errors = []
    rows = []
    for item in items:
        picked = decide(load_events(out / item["name"]), threshold)
        pred = bool(picked)
        if item["violation"] and pred:
            tp += 1
            if item["times"]:
                errors.append(min(abs(e.t - t) for e, _ in picked for t in item["times"]))
        elif item["violation"]:
            fn += 1
        elif pred:
            fp += 1
        else:
            tn += 1
        rows.append((item, picked))
    n = tp + fp + fn + tn
    return {
        "threshold": threshold,
        "accuracy": round((tp + tn) / n, 3),
        "recall": round(tp / (tp + fn), 3) if tp + fn else None,
        "precision": round(tp / (tp + fp), 3) if tp + fp else None,
        "tp": tp, "fp": fp, "fn": fn, "tn": tn,
        "t_err_median": round(statistics.median(errors), 2) if errors else None,
        "t_within_2s": round(sum(e <= 2 for e in errors) / len(errors), 3) if errors else None,
        "rows": rows,
    }


def main() -> None:
    p = argparse.ArgumentParser()
    p.add_argument("--videos", type=Path, default=config.ROOT / "data/open/videos")
    p.add_argument("--labels", type=Path, default=config.ROOT / "data/open/open_labels.xlsx")
    p.add_argument("--out", type=Path, default=config.RUNS_DIR / "eval" / "probe")
    p.add_argument("--only", nargs="*")
    p.add_argument("--force", action="store_true")
    p.add_argument("--parallel", type=int, default=2)
    p.add_argument("--score-only", action="store_true")
    p.add_argument("--candidate", type=float, default=config.CANDIDATE)
    p.add_argument("--analyzer", choices=["probe", "twopass"], default="probe")
    p.add_argument("--scan-from", type=Path)
    args = p.parse_args()

    items = build_items(args.videos, load_labels(args.labels))
    if args.only:
        items = [i for i in items if set(i["aliases"]) & {norm(x) for x in args.only}]
    if not args.score_only:
        asyncio.run(analyze_all(items, args.out, args.force, args.parallel, args.candidate, args.analyzer, args.scan_from))
    items = [i for i in items if (args.out / i["name"] / "raw_events.json").exists()]

    summary = []
    for th in (0.3, 0.5, 0.6, 0.7, 0.8, 0.9):
        s = score(items, args.out, th)
        summary.append({k: v for k, v in s.items() if k != "rows"})
        print({k: v for k, v in s.items() if k != "rows"})

    s = score(items, args.out, config.THRESHOLD)
    print(f"\nthreshold={config.THRESHOLD}")
    for item, picked in s["rows"]:
        label = f"V@{','.join(f'{t:g}' for t in item['times'])}" if item["violation"] else "ok"
        pred = "; ".join(f"{e.rule_id}@{e.t:g}({e.confidence:.2f}x{n})" for e, n in picked) or "-"
        mark = "" if bool(picked) == item["violation"] else "  <-- WRONG"
        print(f"{'/'.join(item['aliases']):>10} {label:>10}  {pred}{mark}")
    (args.out / "summary.json").write_text(json.dumps(summary, indent=1), encoding="utf-8")


if __name__ == "__main__":
    main()
